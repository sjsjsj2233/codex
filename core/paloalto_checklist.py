"""
Palo Alto (PAN-OS) 네트워크 장비 점검 — 작업 전/후 명령어별 원본 출력 비교

체크리스트 항목을 미리 정해두고 값을 요약하는 방식이 아니라, 작업 전/후 파일에
실제로 들어있는 show/request 명령어를 그대로 따라가며 각 명령어의 원본 출력을
줄 단위로 diff한다 (core/config_diff.py의 범용 diff 엔진 재사용). 점검에 사용하는
명령어 목록이 바뀌어도 파일 내용만 보고 자동으로 따라가므로 고정 목록에 의존하지 않는다.
"""
import os
import re
from dataclasses import dataclass, field
from typing import List, Tuple

from core.config_diff import ConfigComparator, DiffType


# ── 점검 명령어 참고 목록 (Cisco의 terminal length 0 대응: set cli pager off) ──
# 실제 비교 로직은 이 목록에 의존하지 않는다 — 장비에서 실행할 명령을 추천하는
# 참고용 안내일 뿐이며, 사용자가 자유롭게 추가/변경해도 비교 결과에는 영향 없다.
PALOALTO_INSPECTION_COMMANDS = (
    "! ===== Palo Alto (PAN-OS) 점검 명령어 =====\n"
    "set cli pager off\n"
    "set cli terminal width 500\n"
    "show system info\n"
    "show session info\n"
    "show running resource-monitor\n"
    "show system resources\n"
    "show system environmentals\n"
    "show system software status\n"
    "show system disk-space\n"
    "show high-availability all\n"
    "show high-availability state-synchronization\n"
    "show arp all\n"
    "show routing route\n"
    "show interface all\n"
    "show interface ethernet1/x\n"
    "request license info\n"
    "show logging-status\n"
    "show jobs all\n"
    "show ntp\n"
    "show admins all\n"
    "show vpn ipsec-sa"
)


# ── 파일 읽기 / 섹션 분리 ────────────────────────────────────────────────────
def _read(path: str) -> str:
    for enc in ('utf-8', 'euc-kr', 'cp949', 'latin-1'):
        try:
            with open(path, 'r', encoding=enc, errors='ignore') as f:
                return f.read()
        except Exception:
            continue
    return ''


# PAN-OS 운영모드 프롬프트: hostname> show ... / username@hostname(active)> show ...
# (실사용 장비는 'sjkim21@8F_NA12_DMZ_FW_PA3220_1_59.148(active)>' 처럼
#  사용자명@ 접두사와 HA 상태(active/passive) 접미사가 붙는 경우가 흔함)
_CMD_RE = re.compile(
    r'(?:^|\n)(?:[\w\-\.]+@)?([\w\-\.]+)(?:\([^)]*\))?[>#]\s*((?:show|request)\s+\S[^\n]*)',
    re.IGNORECASE
)

# 터미널이 다시 그릴 때 섞여 들어오는 ANSI 이스케이프(ESC[K 등) 제거용
_ANSI_ESCAPE_RE = re.compile(r'\x1b(?:\[[0-9;?]*[a-zA-Z]|[=>])')

# "네트워크 자동화" 탭이 만든 TXT 결과 형식: 실제 장비 프롬프트 대신
# 자체적으로 'Command: <cmd>  [Host: ...]' 라벨을 붙여서 저장한다.
_NA_CMD_RE = re.compile(r'(?:^|\n)Command:\s*(.+?)\n', re.IGNORECASE)


def _split_sections(text: str) -> dict:
    # 1) '네트워크 자동화' 탭이 생성한 'Command: <cmd>' 라벨 형식 우선 시도
    na_matches = list(_NA_CMD_RE.finditer(text))
    if na_matches:
        sections = {}
        for i, m in enumerate(na_matches):
            cmd = re.sub(r'\s*\[Host:[^\]]*\]\s*$', '', m.group(1).strip()).strip()
            start = m.end()
            end   = na_matches[i + 1].start() if i + 1 < len(na_matches) else len(text)
            sections.setdefault(cmd, '')
            sections[cmd] += text[start:end]
        return sections

    # 2) 장비에서 직접 복사해온 원본 CLI 프롬프트 형식
    matches = list(_CMD_RE.finditer(text))
    if not matches:
        return {'raw': text}
    sections = {}
    for i, m in enumerate(matches):
        cmd   = m.group(2).strip()
        start = m.end()
        end   = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections.setdefault(cmd, '')
        sections[cmd] += text[start:end]
    return sections


def _extract_hostname(text: str, filepath: str) -> str:
    m = re.search(r'^hostname:\s*(\S+)', text, re.IGNORECASE | re.MULTILINE)
    if m:
        return m.group(1)
    return os.path.splitext(os.path.basename(filepath))[0]


# ── 공개 데이터 모델 ─────────────────────────────────────────────────────────
@dataclass
class CommandDiffResult:
    command: str
    status: str                              # 'changed' | 'unchanged' | 'added' | 'removed' | 'unparsed'
    before_text: str = ''
    after_text: str = ''
    diff_lines: list = field(default_factory=list)   # core.config_diff.DiffLine 목록 (changed일 때만)
    added_count: int = 0
    removed_count: int = 0


# ── 공개 API ──────────────────────────────────────────────────────────────
def compare_files_full(before_path: str, after_path: str) -> Tuple[str, List[CommandDiffResult]]:
    """작업 전/후 파일에 실제로 있는 명령어를 그대로 따라가며 원본 출력을 줄 단위로 비교한다."""
    before_text = _ANSI_ESCAPE_RE.sub('', _read(before_path))
    after_text  = _ANSI_ESCAPE_RE.sub('', _read(after_path))

    hostname = _extract_hostname(after_text, after_path) or _extract_hostname(before_text, before_path)

    before_sections = _split_sections(before_text)
    after_sections  = _split_sections(after_text)

    comparator = ConfigComparator()
    comparator.case_sensitive = True  # 원문 그대로 비교/표시 (대소문자 임의 변경도 실제 차이로 보여줘야 함)

    # 명령어 구간을 아예 못 찾은 경우(둘 다 raw로만 떨어짐) — 파일 전체를 통째로 비교
    if set(before_sections.keys()) == {'raw'} and set(after_sections.keys()) == {'raw'}:
        diff_lines, summary = comparator.compare_strings(before_sections['raw'], after_sections['raw'])
        status = 'changed' if (summary.added_count or summary.removed_count) else 'unchanged'
        return hostname, [CommandDiffResult(
            command='(명령어 구간을 인식하지 못해 파일 전체를 비교했습니다)',
            status=status,
            before_text=before_sections['raw'], after_text=after_sections['raw'],
            diff_lines=diff_lines,
            added_count=summary.added_count, removed_count=summary.removed_count,
        )]

    before_cmds = [c for c in before_sections.keys() if c != 'raw']
    after_cmds  = [c for c in after_sections.keys() if c != 'raw']
    all_cmds = list(dict.fromkeys(before_cmds + after_cmds))  # 순서 보존 합집합

    results = []
    for cmd in all_cmds:
        b = before_sections.get(cmd)
        a = after_sections.get(cmd)
        if b is None:
            results.append(CommandDiffResult(command=cmd, status='added', after_text=a or ''))
            continue
        if a is None:
            results.append(CommandDiffResult(command=cmd, status='removed', before_text=b or ''))
            continue
        diff_lines, summary = comparator.compare_strings(b, a)
        has_changes = bool(summary.added_count or summary.removed_count)
        results.append(CommandDiffResult(
            command=cmd,
            status='changed' if has_changes else 'unchanged',
            before_text=b, after_text=a,
            diff_lines=diff_lines,
            added_count=summary.added_count, removed_count=summary.removed_count,
        ))
    return hostname, results


_STATUS_LABEL = {
    'changed':   '변경됨',
    'unchanged': '변경없음',
    'added':     '신규(작업 후)',
    'removed':   '삭제(작업 후)',
    'unparsed':  '파싱 실패',
}


def _tsv_safe(v: str) -> str:
    return (v or '').replace('\t', ' ').replace('\r', '')


def to_tsv_lines(host_results: List[Tuple[str, List[CommandDiffResult]]]) -> str:
    """엑셀 붙여넣기용 탭구분 텍스트 — 명령어의 각 줄을 작업전/작업후/비교(동일·다름)로
    한 줄씩 펼쳐서 담는다 (사용자가 직접 엑셀 수식으로 만들던 줄 단위 비교 방식과 동일한 구조)."""
    comparator = ConfigComparator()
    comparator.case_sensitive = True

    lines = ['장비\t명령어\t작업전\t작업후\t비교']
    for hostname, results in host_results:
        for r in results:
            if r.status in ('changed', 'unparsed'):
                rows = comparator._make_sidebyside_rows(
                    r.before_text.splitlines(), r.after_text.splitlines()
                )
                for left, right, rtype in rows:
                    same = '동일' if rtype == 'equal' else '다름'
                    lines.append(f'{hostname}\t{r.command}\t{_tsv_safe(left)}\t{_tsv_safe(right)}\t{same}')
            elif r.status == 'unchanged':
                for line in r.before_text.splitlines():
                    lines.append(f'{hostname}\t{r.command}\t{_tsv_safe(line)}\t{_tsv_safe(line)}\t동일')
            elif r.status == 'added':
                for line in r.after_text.splitlines():
                    lines.append(f'{hostname}\t{r.command}\t\t{_tsv_safe(line)}\t다름(신규 명령어)')
            elif r.status == 'removed':
                for line in r.before_text.splitlines():
                    lines.append(f'{hostname}\t{r.command}\t{_tsv_safe(line)}\t\t다름(삭제된 명령어)')
    return '\n'.join(lines)


def save_xlsx_lines(host_results: List[Tuple[str, List[CommandDiffResult]]], output_path: str) -> None:
    """줄 단위 비교 결과를 실제 .xlsx 파일로 저장한다 — '다름' 행은 자동으로 색이 칠해져
    나오므로 엑셀에서 별도 수식/조건부 서식을 넣을 필요가 없다."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def safe(v: str) -> str:
        return ILLEGAL_CHARACTERS_RE.sub('', v or '')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '비교결과'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    diff_fill   = PatternFill(start_color='FECACA', end_color='FECACA', fill_type='solid')
    diff_font   = Font(color='B91C1C', bold=True)
    body_font   = Font(name='Consolas', size=9)

    headers = ['장비', '명령어', '작업전', '작업후', '비교']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['C'].width = 55
    ws.column_dimensions['D'].width = 55
    ws.column_dimensions['E'].width = 16

    comparator = ConfigComparator()
    comparator.case_sensitive = True

    def add_row(hostname, command, before, after, compare_label, is_diff):
        ws.append([hostname, command, safe(before), safe(after), compare_label])
        r = ws.max_row
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = body_font
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (3, 4)))
            if is_diff:
                cell.fill = diff_fill
        if is_diff:
            ws.cell(row=r, column=5).font = diff_font

    for hostname, results in host_results:
        for r in results:
            if r.status in ('changed', 'unparsed'):
                rows = comparator._make_sidebyside_rows(
                    r.before_text.splitlines(), r.after_text.splitlines()
                )
                for left, right, rtype in rows:
                    same = rtype == 'equal'
                    add_row(hostname, r.command, left or '', right or '',
                            '동일' if same else '다름', not same)
            elif r.status == 'unchanged':
                for line in r.before_text.splitlines():
                    add_row(hostname, r.command, line, line, '동일', False)
            elif r.status == 'added':
                for line in r.after_text.splitlines():
                    add_row(hostname, r.command, '', line, '다름(신규 명령어)', True)
            elif r.status == 'removed':
                for line in r.before_text.splitlines():
                    add_row(hostname, r.command, line, '', '다름(삭제된 명령어)', True)

    wb.save(output_path)
